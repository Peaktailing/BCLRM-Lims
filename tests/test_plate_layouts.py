"""Plate layouts (app/plates.py): which of an experiment's animals is in which
well, kept to read the next plate onto; a plate reader's or qPCR machine's
export read in the shapes instruments write, laid on the layout and turned
into each animal's readout (replicates averaged, the blank taken off, ΔCt
from two genes), flagged where a well is off, and saved only when asked,
with the plate kept as the source."""
from __future__ import annotations

import io
import json
import unittest
from datetime import date, timedelta

from tests.base import *  # noqa: F401,F403
from tests.base import AppTestCase, count, location, one, rows, uniq

from app import actions, plates
from app.app import app

TODAY = date.today()

GRID = """Plate reader export
Read 1:450
\t1\t2\t3\t4\t5\t6\t7\t8\t9\t10\t11\t12
A\t0.512\t0.498\t0.505\t0.910\t0.890\t0.930\t\t\t\t\t\t
B\t0.060\t0.050\t\t\t\t\t\t\t\t\t\t
C\t0.210\t0.200\t0.405\t0.395\t\t\t\t\t\t\t\t
"""

QPCR = """* Block Type = 96-Well Block (0.2mL)
* Experiment Name = run 12

[Results]
Well,Well Position,Omit,Sample Name,Target Name,Task,Reporter,CT,Ct Mean
1,A1,false,m1,Gapdh,UNKNOWN,SYBR,18.1,18.15
2,A2,false,m1,Gapdh,UNKNOWN,SYBR,18.2,18.15
3,A3,false,m1,Il6,UNKNOWN,SYBR,25.0,25.1
4,A4,false,m1,Il6,UNKNOWN,SYBR,25.2,25.1
13,B1,false,m2,Gapdh,UNKNOWN,SYBR,18.0,18.0
14,B2,false,m2,Gapdh,UNKNOWN,SYBR,18.0,18.0
15,B3,false,m2,Il6,UNKNOWN,SYBR,24.0,24.0
16,B4,false,m2,Il6,UNKNOWN,SYBR,24.0,24.0
25,C1,false,m3,Gapdh,UNKNOWN,SYBR,18.4,18.4
26,C2,false,m3,Gapdh,UNKNOWN,SYBR,18.4,18.4
27,C3,false,m3,Il6,UNKNOWN,SYBR,23.4,23.4
28,C4,false,m3,Il6,UNKNOWN,SYBR,Undetermined,
37,D1,false,NTC,Gapdh,NTC,SYBR,Undetermined,
"""

VALUE = {"key": "custom", "label": "OD450", "unit": "OD", "kind": "value"}


def subjects(*keys, groups=None):
    return [{"key": k, "label": k.split(":")[1], "group": (groups or {}).get(k, "")} for k in keys]


class ReadingExports(unittest.TestCase):
    def test_a_grid_with_row_letters_and_column_numbers(self):
        found = plates.read_export(text=GRID)
        self.assertEqual((found["shape"], found["kind"], found["format"]), ("grid", "value", 96))
        self.assertEqual(found["targets"], ["Read 1:450"])
        values = found["values"]["Read 1:450"]
        self.assertEqual((values["A1"], values["C4"], len(values)), ("0.512", "0.395", 12))

    def test_two_reads_in_one_file_are_two_targets(self):
        found = plates.read_export(text=GRID + "\n450 nm corrected\n" + GRID.split("\n", 2)[2])
        self.assertEqual(found["targets"], ["Read 1:450", "450 nm corrected"])

    def test_a_qpcr_run_with_several_genes(self):
        found = plates.read_export(text=QPCR)
        self.assertEqual((found["shape"], found["kind"]), ("list", "ct"))
        self.assertEqual(found["targets"], ["Gapdh", "Il6"])
        self.assertEqual(found["values"]["Il6"]["C4"], "Undetermined")
        self.assertEqual(found["values"]["Gapdh"]["A2"], "18.2")      # the Well Position column, not Well's numbers

    def test_lines_of_well_and_value_and_a_well_written_a01(self):
        found = plates.read_export(text="A01\t0.5\nA02\t0.6\nB1\t0.7\n")
        self.assertEqual(found["values"][""], {"A1": "0.5", "A2": "0.6", "B1": "0.7"})
        self.assertEqual(plates.well_name(" h 12"), "H12")
        self.assertIsNone(plates.well_name("Q1"))

    def test_semicolons_and_decimal_commas_and_a_title_line(self):
        found = plates.read_export(filename="od.csv", data="Plate 1\nWell;OD\nA1;0,5\nA2;0,6\n".encode("cp1252"))
        self.assertEqual(found["values"][""], {"A1": "0,5", "A2": "0,6"})
        self.assertEqual(plates.number("0,5"), 0.5)

    def test_a_384_well_export_is_seen_as_one(self):
        found = plates.read_export(text="Well,OD\nA1,0.1\nP24,0.2\n")
        self.assertEqual(found["format"], 384)

    def test_an_excel_workbook(self):
        from openpyxl import Workbook
        book = Workbook()
        book.active.title = "Amplification"
        book.active.append(["Cycle", "Well", "Delta Rn"])
        results = book.create_sheet("Results")
        for line in QPCR.splitlines()[4:]:
            results.append(line.split(","))
        data = io.BytesIO()
        book.save(data)
        found = plates.read_export(filename="run.xlsx", data=data.getvalue())
        self.assertEqual(found["targets"], ["Gapdh", "Il6"])

    def test_nothing_plate_like_is_refused_in_words(self):
        with self.assertRaises(plates.PlateError) as caught:
            plates.read_export(text="name,age\nAnna,3\n")
        self.assertIn("No plate found", str(caught.exception))


class LayingItOnTheLayout(unittest.TestCase):
    layout = {"A1": {"role": "sample", "subject": "mouse:1"}, "A2": {"role": "sample", "subject": "mouse:1"},
              "A3": {"role": "sample", "subject": "mouse:1"},
              "A4": {"role": "sample", "subject": "mouse:2"}, "A5": {"role": "sample", "subject": "mouse:2"},
              "A6": {"role": "sample", "subject": "mouse:2"},
              "B1": {"role": "blank"}, "B2": {"role": "blank"},
              "C1": {"role": "standard", "conc": 1}, "C2": {"role": "standard", "conc": 1},
              "C3": {"role": "standard", "conc": 2}, "C4": {"role": "standard", "conc": 2}}

    def analyse(self, text=GRID, readout=VALUE, options=None, layout=None, **kw):
        return plates.analyse(plates.clean_layout(layout or self.layout, 96), 96, plates.read_export(text=text),
                              readout, subjects("mouse:1", "mouse:2"), options, **kw)

    def by(self, result):
        return {r["key"]: r for r in result["subjects"]}

    def test_replicates_averaged_and_the_blank_taken_off(self):
        out = self.analyse()
        self.assertEqual(out["blank"], {"mean": 0.055, "n": 2, "subtracted": True})
        one_, two = self.by(out)["mouse:1"], self.by(out)["mouse:2"]
        self.assertEqual(one_["wells"], ["A1", "A2", "A3"])
        self.assertEqual(one_["reads"]["Read 1:450"]["mean"], 0.505)
        self.assertEqual((one_["value"], two["value"]), (0.45, 0.855))
        self.assertEqual(out["ready"], 2)
        # Without the blank, the mean as it is.
        self.assertEqual(self.by(self.analyse(options={"subtract_blank": False}))["mouse:1"]["value"], 0.505)

    def test_a_replicate_far_from_the_others_is_flagged_and_can_be_left_out(self):
        text = GRID.replace("0.498", "0.950")
        out = self.analyse(text=text)
        self.assertEqual(out["flagged"]["outlier"], ["A2"])
        self.assertEqual(out["wells"]["A2"]["reads"]["Read 1:450"]["flags"], ["outlier"])
        self.assertNotEqual(self.by(out)["mouse:1"]["value"], 0.45)
        left = self.analyse(text=text, options={"skip": ["a2"]})
        self.assertEqual(left["skipped"], ["A2"])
        self.assertTrue(left["wells"]["A2"]["skipped"])
        self.assertEqual(self.by(left)["mouse:1"]["reads"]["Read 1:450"]["n"], 2)
        self.assertEqual(self.by(left)["mouse:1"]["value"], 0.4535)

    def test_a_well_past_what_a_reader_reads_is_left_out(self):
        text = GRID.replace("0.910", "OVRFLW").replace("0.890", "12.5")
        out = self.analyse(text=text, readout={"key": "derived:blank_subtracted", "formula": "blank_subtracted",
                                                 "label": "Absorbance, blank subtracted", "unit": "OD"})
        self.assertEqual(out["flagged"]["limit"], ["A4", "A5"])
        two = self.by(out)["mouse:2"]
        self.assertEqual(two["inputs"], {"od": 0.93, "blank": 0.055})
        self.assertEqual(two["value"], 0.875)

    def test_blank_subtracted_needs_its_blank(self):
        layout = {w: s for w, s in self.layout.items() if s["role"] != "blank"}
        with self.assertRaises(plates.PlateError) as caught:
            self.analyse(layout=layout, readout={"formula": "blank_subtracted", "label": "Absorbance, blank subtracted"})
        self.assertIn("Mark the blank wells", str(caught.exception))

    def test_a_standard_curve(self):
        out = self.analyse(options={"standard_curve": True})
        self.assertEqual([s["conc"] for s in out["standards"]], [1.0, 2.0])
        self.assertAlmostEqual(out["fit"]["slope"], 0.195, places=3)
        self.assertAlmostEqual(out["fit"]["r2"], 1.0, places=2)
        one_ = self.by(out)["mouse:1"]
        self.assertAlmostEqual(one_["value"], (0.45 - out["fit"]["intercept"]) / out["fit"]["slope"], places=3)
        self.assertTrue(self.by(out)["mouse:2"]["outside"])        # read past the highest standard

    def test_a_qpcr_run_gives_delta_ct_from_two_genes(self):
        layout = {"A1": "mouse:1", "A2": "mouse:1", "A3": "mouse:1", "A4": "mouse:1",
                  "B1": "mouse:2", "B2": "mouse:2", "B3": "mouse:2", "B4": "mouse:2",
                  "C1": "mouse:3", "C2": "mouse:3", "C3": "mouse:3", "C4": "mouse:3", "D1": "blank"}
        readout = {"key": "derived:delta_ct", "formula": "delta_ct", "label": "ΔCt"}
        out = plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=QPCR), readout,
                             subjects("mouse:1", "mouse:2", "mouse:3"), {"target": "Il6", "reference": "Gapdh"})
        got = self.by(out)
        self.assertEqual(got["mouse:1"]["inputs"], {"ct_target": 25.1, "ct_reference": 18.15})
        self.assertEqual((got["mouse:1"]["value"], got["mouse:2"]["value"]), (6.95, 6.0))
        self.assertEqual(out["flagged"]["no_value"], ["C4", "D1"])
        self.assertEqual(got["mouse:3"]["value"], 5.0)               # the Il6 well that came up
        # The genes the other way round, or the same twice.
        self.assertEqual(self.by(plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=QPCR),
                                                readout, subjects("mouse:1"), {"target": "Gapdh",
                                                                                "reference": "Il6"}))["mouse:1"]["value"],
                         -6.95)
        with self.assertRaises(plates.PlateError):
            plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=QPCR), readout,
                           subjects("mouse:1"), {"target": "Il6", "reference": "Il6"})

    def test_a_ct_replicate_half_a_cycle_out_is_flagged(self):
        text = QPCR.replace("2,A2,false,m1,Gapdh,UNKNOWN,SYBR,18.2", "2,A2,false,m1,Gapdh,UNKNOWN,SYBR,19.9")
        layout = {"A1": "mouse:1", "A2": "mouse:1", "A3": "mouse:1", "A4": "mouse:1"}
        out = plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=text),
                             {"formula": "delta_ct", "label": "ΔCt"}, subjects("mouse:1"))
        self.assertEqual(sorted(out["flagged"]["outlier"]), ["A1", "A2"])   # two replicates 1.8 cycles apart

    def test_a_readout_a_plate_cant_give_is_refused(self):
        for readout in ({"key": "body_weight", "label": "Body weight", "kind": "value"},
                        {"key": "survival", "label": "Survival", "kind": "fraction"},
                        {"formula": "tumour_lw", "label": "Tumour volume (L × W² / 2)"}):
            with self.assertRaises(plates.PlateError):
                self.analyse(readout=readout)

    def test_a_layout_names_only_wells_on_the_plate_and_animals_in_the_experiment(self):
        with self.assertRaises(plates.PlateError):
            plates.clean_layout({"I1": "mouse:1"}, 96)
        with self.assertRaises(plates.PlateError):
            plates.clean_layout({"A1": "mouse:9"}, 96, {"mouse:1"})
        with self.assertRaises(plates.PlateError):
            plates.clean_layout({"A1": {"role": "standard"}}, 96)
        self.assertEqual(plates.clean_layout({"a01": "blank", "A2": "", "P24": {"role": "standard", "conc": "2,5"}}, 384),
                         {"A1": {"role": "blank"}, "P24": {"role": "standard", "conc": 2.5}})


def elisa(*samples, a=0.05, b=1.2, c=150.0, d=3.0, noise=None):
    """An ELISA plate as "well value" lines: eight standards in duplicate
    (rows C and D, 0 to 1000) on a 4PL curve, then each sample's two wells
    along row A at the concentration given (or, as a string, that signal)."""
    from app.standard_curves import logistic
    concs = [0, 15.625, 31.25, 62.5, 125, 250, 500, 1000]
    lines, layout = [], {}
    for i, conc in enumerate(concs):
        for row in "CD":
            y = logistic(conc, a, b, c, d) * (1 + (noise() if noise else 0))
            lines.append(f"{row}{i + 1}\t{y:.5f}")
            layout[f"{row}{i + 1}"] = {"role": "standard", "conc": conc}
    for i, (key, conc) in enumerate(samples):
        y = conc if isinstance(conc, str) else f"{logistic(conc, a, b, c, d):.5f}"
        for col in (2 * i + 1, 2 * i + 2):
            lines.append(f"A{col}\t{y}")
            layout[f"A{col}"] = {"role": "sample", "subject": key}
    return "\n".join(lines), layout


class StandardCurves(unittest.TestCase):
    """app/standard_curves.py: a line, a log–log line, a 4PL by
    Levenberg–Marquardt in pure Python, refused in words when it doesn't settle."""

    def points(self, a=0.05, b=1.2, c=150.0, d=3.0, concs=(0, 15.625, 31.25, 62.5, 125, 250, 500, 1000), noise=None):
        from app.standard_curves import logistic
        return [(x, logistic(x, a, b, c, d) * (1 + (noise() if noise else 0))) for x in concs for _ in range(2)]

    def test_a_4pl_finds_the_curve_it_was_drawn_from(self):
        from app import standard_curves as sc
        for a, b, c, d in ((0.05, 1.2, 150.0, 3.0), (2.5, 0.8, 40.0, 0.1), (0.1, 2.5, 300.0, 1.8)):   # rising, falling, steep
            f = sc.four_pl(self.points(a, b, c, d))
            for got, want in zip((f["a"], f["b"], f["c"], f["d"]), (a, b, c, d)):
                self.assertAlmostEqual(got, want, delta=abs(want) * 1e-3)
            self.assertEqual(f["r2"], 1.0)
            self.assertAlmostEqual(sc.concentration(f, sc.logistic(200, a, b, c, d)), 200, delta=0.01)

    def test_a_4pl_through_noisy_standards(self):
        import random
        from app import standard_curves as sc
        rng = random.Random(7)
        f = sc.four_pl(self.points(noise=lambda: rng.gauss(0, 0.03)))
        self.assertGreater(f["r2"], 0.98)
        self.assertAlmostEqual(f["c"], 150, delta=150 * 0.2)
        self.assertAlmostEqual(f["d"], 3.0, delta=0.3)
        self.assertAlmostEqual(f["b"], 1.2, delta=0.25)
        self.assertGreater(f["rmse"], 0)
        self.assertAlmostEqual(sc.concentration(f, sc.logistic(100, 0.05, 1.2, 150, 3.0)), 100, delta=15)

    def test_a_signal_the_curve_never_reaches(self):
        from app import standard_curves as sc
        f = sc.four_pl(self.points())
        self.assertIsNone(sc.concentration(f, 3.5))          # past the top plateau
        self.assertFalse(sc.below(f, 3.5))
        self.assertIsNone(sc.concentration(f, 0.01))         # under the bottom one
        self.assertTrue(sc.below(f, 0.01))

    def test_a_4pl_that_does_not_settle_is_refused_in_words(self):
        from app import standard_curves as sc
        for points in ([(x, 0.5) for x in (1, 2, 4, 8, 16)],                         # flat
                       [(1, 1), (2, 0.2), (4, 0.1), (8, 0.2), (16, 1)],                # a V, not an S
                       [(1, 1), (10, 0.1), (100, 1), (1000, 0.1), (10000, 1)]):        # a zigzag
            with self.assertRaises(sc.CurveError) as caught:
                sc.four_pl(points)
            self.assertTrue(str(caught.exception))
        with self.assertRaises(sc.CurveError) as caught:
            sc.four_pl([(1, 0.1), (2, 0.2), (4, 0.4), (8, 0.8)])
        self.assertIn("at least five concentrations", str(caught.exception))

    def test_log_log_and_linear(self):
        from app import standard_curves as sc
        f = sc.loglog([(0, 0.0)] + [(x, 0.01 * x ** 0.8) for x in (10, 100, 1000)])
        self.assertEqual((f["slope"], f["intercept"], f["r2"], f["left_out"]), (0.8, -2.0, 1.0, 1))
        self.assertAlmostEqual(sc.concentration(f, 0.01 * 50 ** 0.8), 50, places=6)
        self.assertIsNone(sc.concentration(f, 0))
        line = sc.linear([(1, 0.25), (2, 0.45), (4, 0.85)])
        self.assertEqual((line["slope"], line["intercept"]), (0.2, 0.05))
        with self.assertRaises(sc.CurveError):
            sc.loglog([(0, 0.1), (5, 0.2)])

    def test_the_curve_chosen_for_the_standards(self):
        from app import standard_curves as sc
        self.assertEqual(sc.fit(self.points())["model"], "4pl")            # eight, over a 64-fold range
        self.assertTrue(sc.fit(self.points())["auto"])
        self.assertEqual(sc.fit(self.points(concs=(1, 2, 3)))["model"], "linear")
        self.assertEqual(sc.fit(self.points(concs=(10, 11, 12, 13, 14)))["model"], "linear")   # not a tenfold range
        self.assertEqual(sc.fit([(1, 1), (10, 0.1), (100, 1), (1000, 0.1), (10000, 1)])["model"], "linear")
        self.assertEqual(sc.fit(self.points(), "loglog")["model"], "loglog")
        self.assertNotIn("auto", sc.fit(self.points(), "linear"))


class AFourPLPlate(unittest.TestCase):
    def test_each_animal_is_read_off_the_curve_and_the_curve_is_there_to_draw(self):
        text, layout = elisa(("mouse:1", 100), ("mouse:2", 2000), ("mouse:3", "3.5000"), ("mouse:4", "0.0100"))
        out = plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=text), VALUE,
                             subjects("mouse:1", "mouse:2", "mouse:3", "mouse:4"), {"standard_curve": True})
        self.assertEqual((out["fit"]["model"], out["fit"]["auto"]), ("4pl", True))
        got = {r["key"]: r for r in out["subjects"]}
        self.assertAlmostEqual(got["mouse:1"]["value"], 100, delta=0.5)
        self.assertNotIn("outside", got["mouse:1"])
        self.assertAlmostEqual(got["mouse:2"]["value"], 2000, delta=20)
        self.assertTrue(got["mouse:2"]["outside"])                    # past the highest standard, still marked
        self.assertIn("dilute", got["mouse:3"]["error"])
        self.assertIn("less than the assay detects", got["mouse:4"]["error"])
        self.assertEqual(out["ready"], 2)
        curve = out["curve"]
        self.assertTrue(curve["log"])
        self.assertEqual(len(curve["standards"]), 16)
        self.assertGreater(len(curve["line"]), 20)
        self.assertEqual([round(x) for x, _y in curve["samples"]], [100, 2000])
        # Asked for a line instead, it is a line.
        line = plates.analyse(plates.clean_layout(layout, 96), 96, plates.read_export(text=text), VALUE,
                              subjects("mouse:1"), plates.options_from({"standard_curve": "1", "curve": "linear"}))
        self.assertEqual(line["fit"]["model"], "linear")
        self.assertNotIn("auto", line["fit"])
        self.assertEqual(plates.options_from({"curve": "Log–log"})["curve"], "loglog")


class Base(AppTestCase):
    def experiment(self, groups=("Vehicle", "Vehicle", "Drug"), readout=None, client=None):
        client = client or self.m
        r = client.post("/colony/experiments/create", data={"name": uniq("ELISA "),
                                                            "start_date": (TODAY - timedelta(days=3)).isoformat()})
        exp = int(location(r).rsplit("/", 1)[1])
        keys = []
        for group in groups:
            mouse = self.make_mouse(client, self.member)
            client.post(f"/colony/experiments/{exp}/add-mouse", data={"mouse_row_id": mouse, "treatment_group": group})
            keys.append(f"mouse:{mouse}")
        form = readout or {"readout": "custom", "readout_label": "IL-6", "readout_unit": "pg/mL", "readout_kind": "value"}
        r = client.post(f"/experiments/{exp}/update", data=form, headers={"X-Autosave": "1"})
        self.assertEqual(r.status_code, 200, r.get_data(as_text=True)[:300])
        return exp, keys

    def save_layout(self, exp, wells, name="Plate 1", fmt=96, client=None, **extra):
        return (client or self.m).post(f"/experiments/{exp}/plates/layouts",
                                       json={"name": name, "format": fmt, "wells": wells, **extra})

    def grid_layout(self, keys):
        """Each animal in duplicate along row A, two blanks in B."""
        wells = {}
        for i, key in enumerate(keys):
            wells[f"A{2 * i + 1}"] = wells[f"A{2 * i + 2}"] = {"role": "sample", "subject": key}
        wells["B1"] = wells["B2"] = {"role": "blank"}
        return wells


class Layouts(Base):
    def test_a_layout_is_saved_listed_and_reused_for_the_next_plate(self):
        exp, keys = self.experiment()
        r = self.save_layout(exp, self.grid_layout(keys))
        self.assertEqual(r.status_code, 200, r.get_json())
        first = r.get_json()["id"]
        self.assertEqual(json.loads(one("select wells from plate_layouts where id=?", first))["B1"], {"role": "blank"})
        other = self.save_layout(exp, {"A1": keys[0]}, name="Plate 2").get_json()["id"]
        # Same as last plate: the layout last read onto comes first.
        self.read_plate(exp, first)
        listed = self.m.get(f"/experiments/{exp}/plates.json").get_json()
        self.assertEqual([x["id"] for x in listed["layouts"]], [first, other])
        self.assertIsNotNone(listed["layouts"][0]["used_at"])
        # Renamed in place; a second layout of the same name is refused.
        self.assertEqual(self.save_layout(exp, self.grid_layout(keys), name="ELISA", id=first).get_json()["id"], first)
        clash = self.save_layout(exp, {}, name="elisa")
        self.assertEqual(clash.status_code, 409)
        self.assertIn("already has a layout", clash.get_json()["error"])
        # An animal not in the experiment, a well off the plate.
        self.assertIn("isn't in this experiment", self.save_layout(exp, {"A1": "mouse:999999"}, name="x").get_json()["error"])
        self.assertIn("isn't a well", self.save_layout(exp, {"E1": keys[0]}, name="x", fmt=24).get_json()["error"])
        # Deleting one keeps the plates read onto it, with their copy.
        self.assertEqual(self.m.post(f"/experiments/{exp}/plates/layouts/{first}/delete").status_code, 200)
        self.assertEqual(count("plate_reads", "experiment_id_fk=? and layout_id_fk is null", exp), 1)

    def test_deleting_the_experiment_takes_its_layouts_and_plates(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        self.read_plate(exp, layout)
        self.assertEqual(self.m.post(f"/experiments/{exp}/delete").status_code, 302)
        self.assertEqual((count("plate_layouts", "experiment_id_fk=?", exp), count("plate_reads", "experiment_id_fk=?", exp)),
                         (0, 0))

    def read_plate(self, exp, layout_id, text=GRID, **form):
        r = self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout_id, "text": text, **form})
        self.assertEqual(r.status_code, 200, r.get_json())
        return r.get_json()

    def test_only_who_may_edit_the_experiment_lays_out_and_reads_plates(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        self.assertEqual(self.save_layout(exp, {}, name="Mine", client=self.o).status_code, 403)
        self.assertEqual(self.o.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID}).status_code, 403)
        self.assertEqual(self.o.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID}).status_code, 403)
        self.assertEqual(self.o.post(f"/experiments/{exp}/plates/layouts/{layout}/delete").status_code, 403)
        self.assertEqual(count("experiment_readings", "experiment_id_fk=?", exp), 0)
        # Seeing them is for anyone who sees the experiment; an admin may read the plate.
        self.assertEqual(len(self.o.get(f"/experiments/{exp}/plates.json").get_json()["layouts"]), 1)
        self.assertEqual(self.a.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID}).status_code, 200)

    def test_the_page_offers_the_plate_to_its_owner_only(self):
        exp, _keys = self.experiment()
        self.assertIn("data-plate-open", self.get_ok(self.m, f"/colony/experiments/{exp}"))
        self.assertIn("plate-layout.js", self.get_ok(self.m, f"/colony/experiments/{exp}"))
        self.assertNotIn("data-plate-open", self.get_ok(self.o, f"/colony/experiments/{exp}"))


class ReadingAPlate(Base):
    def test_preview_changes_nothing_and_save_writes_each_animals_reading(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        r = self.m.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID})
        self.assertEqual(r.status_code, 200, r.get_json())
        out = r.get_json()
        self.assertEqual(out["ready"], 3)
        self.assertEqual([s["value"] for s in out["subjects"]], [0.45, 0.6525, 0.855])
        self.assertEqual(count("experiment_readings", "experiment_id_fk=?", exp), 0)
        self.assertEqual(count("plate_reads", "experiment_id_fk=?", exp), 0)
        on = (TODAY - timedelta(days=1)).isoformat()
        saved = self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "on": on})
        self.assertEqual(saved.status_code, 200, saved.get_json())
        self.assertEqual(saved.get_json()["saved"], 3)
        self.assertIn("Saved 3 readings", saved.get_json()["message"])
        got = dict(rows("select subject, value from experiment_readings where experiment_id_fk=? and read_on=?", exp, on))
        self.assertEqual(got, {keys[0]: 0.45, keys[1]: 0.6525, keys[2]: 0.855})
        read = rows("select layout_name, readout_key, results, data, file_path from plate_reads where experiment_id_fk=?", exp)[0]
        self.assertEqual(read[:2], ("Plate 1", "custom"))
        self.assertEqual(json.loads(read[2])[keys[1]], 0.6525)
        self.assertEqual(json.loads(read[3])["Read 1:450"]["A1"], "0.512")
        self.assertEqual(read[4], "")
        # The sheet shows them; the same plate again replaces, saying so first.
        table = saved.get_json()["table"]
        self.assertEqual(next(x for x in table["rows"] if x["key"] == keys[0])["values"], [0.45])
        again = self.m.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID, "on": on})
        self.assertEqual(again.get_json()["replaces"], 3)

    def test_a_file_is_kept_as_the_source_and_served_as_the_experiment_is(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        r = self.m.post(f"/experiments/{exp}/plates/save", content_type="multipart/form-data",
                        data={"layout_id": str(layout), "file": (io.BytesIO(GRID.encode("utf-8")), "elisa.tsv")})
        self.assertEqual(r.status_code, 200, r.get_json())
        path, name = rows("select file_path, file_name from plate_reads where experiment_id_fk=?", exp)[0]
        self.assertTrue(path.startswith("uploads/"))
        self.assertEqual(name, "elisa.tsv")
        listed = self.m.get(f"/experiments/{exp}/plates.json").get_json()["reads"][0]
        self.assertEqual((listed["count"], listed["file_url"]), (3, f"/static/{path}"))
        served = self.o.get(listed["file_url"])
        self.assertEqual(served.status_code, 200)
        served.close()

    def test_nothing_is_kept_when_nothing_can_be_saved(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, {"A1": keys[0], "B1": "blank"}).get_json()["id"]
        r = self.m.post(f"/experiments/{exp}/plates/save", content_type="multipart/form-data",
                        data={"layout_id": str(layout), "file": (io.BytesIO(b"A1\t0.01\nB1\t0.5\n"), "low.tsv")})
        self.assertEqual(r.status_code, 400)
        self.assertIn("below zero", r.get_json()["error"])
        self.assertEqual(count("plate_reads", "experiment_id_fk=?", exp), 0)
        refused = self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID,
                                                                       "on": (TODAY + timedelta(days=2)).isoformat()})
        self.assertIn("still to come", refused.get_json()["error"])

    def test_fold_change_from_a_qpcr_run_against_the_control_group(self):
        exp, keys = self.experiment(groups=("Vehicle", "Drug", "Drug"),
                                    readout={"readout": "derived:fold_change", "readout_control": "Vehicle"})
        wells = {}
        for row, key in zip("ABC", keys):
            for col in range(1, 5):
                wells[f"{row}{col}"] = key
        layout = self.save_layout(exp, wells, name="qPCR").get_json()["id"]
        form = {"layout_id": layout, "text": QPCR, "target": "Il6", "reference": "Gapdh"}
        out = self.m.post(f"/experiments/{exp}/plates/preview", data=form).get_json()
        values = {s["key"]: s["value"] for s in out["subjects"]}
        self.assertEqual(values[keys[0]], 1.0)                   # the control against its own mean
        self.assertAlmostEqual(values[keys[1]], 2 ** 0.95, places=3)
        self.assertEqual(self.m.post(f"/experiments/{exp}/plates/save", data=form).get_json()["saved"], 3)
        note = one("select note from experiment_readings where experiment_id_fk=? and subject=?", exp, keys[0])
        self.assertEqual(json.loads(note), {"ct_target": 25.1, "ct_reference": 18.15})


class AnotherMeasure(Base):
    """A plate fills the main readout, or another measure of the same
    animals beside it: kept with the experiment, on the page, in Export,
    in experiment_summary."""

    def tumour_experiment(self):
        return self.experiment(readout={"readout": "tumour_volume"})

    def test_a_plate_fills_another_measure_beside_the_main_readout(self):
        exp, keys = self.tumour_experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        listed = self.m.get(f"/experiments/{exp}/plates.json").get_json()
        self.assertEqual((listed["main"]["key"], listed["main"]["plate"], listed["others"]), ("tumour_volume", True, []))
        form = {"layout_id": layout, "text": GRID, "measure": "new", "measure_label": "IL-6 ELISA",
                "measure_unit": "pg/mL"}
        # Trying it changes nothing, the measure included.
        preview = self.m.post(f"/experiments/{exp}/plates/preview", data=form).get_json()
        self.assertEqual((preview["readout"]["label"], preview["readout"]["measure"]), ("IL-6 ELISA", True))
        self.assertNotIn("measures", one("select readout from experiments where id=?", exp))
        saved = self.m.post(f"/experiments/{exp}/plates/save", data=form).get_json()
        self.assertEqual((saved["saved"], saved["filled"]), (3, "measure:1"), saved)
        self.assertIn("IL-6 ELISA", saved["message"])
        self.assertEqual(count("experiment_readings", "experiment_id_fk=? and readout_key=?", exp, "measure:1"), 3)
        self.assertEqual(count("experiment_readings", "experiment_id_fk=? and readout_key=?", exp, "tumour_volume"), 0)
        # The page: the main readout's sheet is as it was, the measure has its own.
        self.assertEqual(saved["table"]["dates"], [])
        measure = saved["measures"][0]
        self.assertEqual((measure["key"], measure["label"], measure["unit"]), ("measure:1", "IL-6 ELISA", "pg/mL"))
        self.assertEqual(next(r for r in measure["table"]["rows"] if r["key"] == keys[0])["values"], [0.45])
        self.assertIn("xp-measure-chips", self.get_ok(self.m, f"/colony/experiments/{exp}"))
        listed = self.m.get(f"/experiments/{exp}/plates.json").get_json()
        self.assertEqual(listed["others"], [{"key": "measure:1", "label": "IL-6 ELISA", "unit": "pg/mL"}])
        self.assertEqual(listed["reads"][0]["readout"], "IL-6 ELISA")
        # The next plate picks it by name or key; no second measure of that name.
        again = self.m.post(f"/experiments/{exp}/plates/save", data={**form, "measure_label": "il-6 elisa"}).get_json()
        self.assertEqual(again["filled"], "measure:1")
        self.assertEqual(len(again["measures"]), 1)
        self.assertEqual(self.m.post(f"/experiments/{exp}/plates/preview",
                                     data={"layout_id": layout, "text": GRID, "measure": "measure:1"}).get_json()["replaces"], 3)
        refused = self.m.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID, "measure": "IFN"})
        self.assertIn("IL-6 ELISA", refused.get_json()["error"])
        self.assertIn("main readout", self.m.post(f"/experiments/{exp}/plates/preview", data={
            **form, "measure_label": "Tumour volume"}).get_json()["error"])
        # A cell typed in its sheet is that measure's; a new main readout keeps it.
        r = self.m.post(f"/experiments/{exp}/readings", json={"on": TODAY.isoformat(), "values": {keys[0]: "0.5"},
                                                              "measure": "measure:1"})
        self.assertEqual(r.status_code, 200, r.get_json())
        self.assertEqual(one("select value from experiment_readings where subject=? and readout_key='measure:1'", keys[0]), 0.5)
        self.assertEqual(self.m.post(f"/experiments/{exp}/readings", json={"on": TODAY.isoformat(), "values": {keys[0]: "1"},
                                                                           "measure": "nope"}).status_code, 400)
        self.m.post(f"/experiments/{exp}/update", data={"readout": "score"}, headers={"X-Autosave": "1"})
        self.assertEqual(json.loads(one("select readout from experiments where id=?", exp))["key"], "score")
        self.assertEqual([m["label"] for m in self.m.get(f"/experiments/{exp}/data.json").get_json()["measures"]],
                         ["IL-6 ELISA"])

    def test_a_main_readout_no_plate_gives_still_has_its_measures(self):
        exp, keys = self.experiment(readout={"readout": "body_weight"})
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        self.assertFalse(self.m.get(f"/experiments/{exp}/plates.json").get_json()["main"]["plate"])
        refused = self.m.post(f"/experiments/{exp}/plates/preview", data={"layout_id": layout, "text": GRID})
        self.assertIn("Another measure", refused.get_json()["error"])
        saved = self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "measure": "new",
                                                                     "measure_label": "Serum ALT"}).get_json()
        self.assertEqual(saved["saved"], 3, saved)
        self.assertEqual(count("mouse_weights", "notes like ?", "ELISA%"), 0)

    def test_export_and_the_summary_carry_it(self):
        exp, keys = self.tumour_experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "measure": "new",
                                                             "measure_label": "IL-6 ELISA", "measure_unit": "pg/mL"})
        r = self.m.get(f"/experiments/{exp}/export.xlsx")
        from openpyxl import load_workbook
        book = load_workbook(io.BytesIO(r.data))
        self.assertIn("IL-6 ELISA", book.sheetnames)
        sheet = list(book["IL-6 ELISA"].values)
        self.assertIn("IL-6 ELISA (pg/mL)", sheet[0])
        self.assertEqual(len(sheet), 4)
        self.assertEqual(len(list(book["Readout"].values)), 1)        # the main readout: nothing yet
        self.assertIn(("IL-6 ELISA",), [tuple(v for v in row if v is not None) for row in book["Tests by day"].values])
        from app.assistant_compute import experiment_summary
        from app.db import SessionLocal
        with app.test_request_context("/"), SessionLocal() as s:
            summary = experiment_summary(s, exp)
        other = summary["readouts"]["other_measures"][0]
        self.assertEqual((summary["readouts"]["label"], other["label"], other["unit"], other["days_measured"]),
                         ("Tumour volume", "IL-6 ELISA", "pg/mL", 1))
        outcome = next(i for i in summary["arrive"] if i["item"] == 6)
        self.assertTrue(any("Secondary measures: IL-6 ELISA" in r for r in outcome["recorded"]))


class UndoingAPlate(Base):
    def test_batch_history_undo_takes_the_whole_plate_back(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        on = (TODAY - timedelta(days=1)).isoformat()
        # One animal had a value that day already: the plate replaces it, and Undo puts it back.
        self.m.post(f"/experiments/{exp}/readings", json={"on": on, "values": {keys[0]: "0.3"}})
        r = self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "on": on})
        self.assertEqual(r.status_code, 200, r.get_json())
        batch, description = rows("select id, description from batches where actor=? order by id desc", self.member)[0]
        self.assertIn("Plate 1", description)
        self.assertEqual(count("audit_log", "batch_id_fk=? and table_name='experiment_readings'", batch), 3)
        self.assertEqual(count("audit_log", "batch_id_fk=? and table_name='plate_reads'", batch), 1)
        self.post(self.m, f"/batches/{batch}/undo")
        self.assertEqual(count("plate_reads", "experiment_id_fk=?", exp), 0)
        self.assertEqual(dict(rows("select subject, value from experiment_readings where experiment_id_fk=?", exp)),
                         {keys[0]: 0.3})
        self.assertEqual(count("plate_layouts", "id=?", layout), 1)        # the layout stays

    def test_a_later_plate_on_the_same_layout_does_not_stop_the_undo(self):
        exp, keys = self.experiment()
        layout = self.save_layout(exp, self.grid_layout(keys)).get_json()["id"]
        first = (TODAY - timedelta(days=2)).isoformat()
        self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "on": first,
                                                             "measure": "new", "measure_label": "IL-6 ELISA"})
        batch = one("select max(id) from batches where actor=?", self.member)
        self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "on": TODAY.isoformat(),
                                                             "measure": "measure:1"})
        self.post(self.m, f"/batches/{batch}/undo")
        self.assertEqual(count("plate_reads", "experiment_id_fk=?", exp), 1)
        self.assertEqual(count("experiment_readings", "experiment_id_fk=? and read_on=?", exp, first), 0)
        self.assertEqual(count("experiment_readings", "experiment_id_fk=? and read_on=?", exp, TODAY.isoformat()), 3)
        # The measure stays, with the later plate's readings to show.
        self.assertIn("IL-6 ELISA", one("select readout from experiments where id=?", exp))


class ForAnAssistant(Base):
    def setUp(self):
        super().setUp()
        from tests.test_api import Base as ApiBase
        from tests.test_mcp import TestApi, bm
        self.bm = bm
        self.api = TestApi(ApiBase.token(self, self.m, scope="propose"))

    def test_the_tool_lays_an_export_on_a_layout_and_gives_the_change_to_propose(self):
        exp, keys = self.experiment()
        self.save_layout(exp, self.grid_layout(keys), name="ELISA 1")
        out = self.bm.call_tool(self.api, "plate_readout", {"experiment_id": exp, "layout": "elisa 1", "text": GRID})
        self.assertTrue(out["changes_nothing"])
        self.assertEqual(out["change"]["fields"]["values"], {keys[0]: 0.45, keys[1]: 0.6525, keys[2]: 0.855})
        self.assertEqual(count("experiment_readings", "experiment_id_fk=?", exp), 0)
        preview = actions.run(app, self.member, [out["change"]])
        self.assertTrue(preview.ok, preview.changes[0].as_dict())
        wrong = self.bm.call_tool(self.api, "plate_readout", {"experiment_id": exp, "layout": "nope", "text": GRID})
        self.assertIn("ELISA 1", wrong["error"])
        inline = self.bm.call_tool(self.api, "plate_readout", {"experiment_id": exp, "format": 96, "text": GRID,
                                                               "wells": {"A1": keys[0], "A2": keys[0]}})
        self.assertEqual(inline["change"]["fields"]["values"], {keys[0]: 0.505})

    def test_the_tool_fills_another_measure_by_its_name(self):
        exp, keys = self.experiment(readout={"readout": "tumour_volume"})
        layout = self.save_layout(exp, self.grid_layout(keys), name="ELISA 1").get_json()["id"]
        self.m.post(f"/experiments/{exp}/plates/save", data={"layout_id": layout, "text": GRID, "measure": "new",
                                                             "measure_label": "IL-6 ELISA",
                                                             "on": (TODAY - timedelta(days=1)).isoformat()})
        out = self.bm.call_tool(self.api, "plate_readout", {"experiment_id": exp, "layout": "ELISA 1", "text": GRID,
                                                            "measure": "IL-6 ELISA"})
        self.assertEqual(out["change"]["fields"]["measure"], "IL-6 ELISA")
        preview = actions.run(app, self.member, [out["change"]])
        self.assertTrue(preview.ok, preview.changes[0].as_dict())
        self.assertIn("IL-6 ELISA", preview.changes[0].as_dict()["summary"])
        self.assertIn("no measure", self.bm.call_tool(self.api, "plate_readout", {
            "experiment_id": exp, "layout": "ELISA 1", "text": GRID, "measure": "IFN"})["error"])
        self.assertIn("Plate layout dialog", self.bm.call_tool(self.api, "plate_readout", {
            "experiment_id": exp, "layout": "ELISA 1", "text": GRID, "measure": "new"})["error"])

    def test_a_captures_file_is_the_export(self):
        exp, keys = self.experiment()
        self.save_layout(exp, self.grid_layout(keys), name="ELISA 1")
        r = self.m.post("/captures/new", content_type="multipart/form-data",
                        data={"text": "plate", "file": (io.BytesIO(GRID.encode("utf-8")), "elisa.csv")})
        self.assertEqual(r.status_code, 302)
        capture = one("select max(id) from captures where owner_username=?", self.member)
        out = self.bm.call_tool(self.api, "plate_readout", {"experiment_id": exp, "layout": "ELISA 1", "capture": capture})
        self.assertEqual(out.get("ready"), 3, out.get("error"))
        # Someone else's capture is not there for them.
        from tests.test_api import Base as ApiBase
        from tests.test_mcp import TestApi
        theirs = TestApi(ApiBase.token(self, self.o, scope="propose"))
        self.assertIn("no capture", self.bm.call_tool(theirs, "plate_readout", {"experiment_id": exp, "layout": "ELISA 1",
                                                                                 "capture": capture})["error"])

    def test_it_is_a_read_only_tool_everywhere(self):
        from pathlib import Path
        from app import assistant_tools, mcp_http
        self.assertIn("plate_readout", assistant_tools.TOOLS)
        self.assertIn("plate_readout", mcp_http.READ_ONLY)
        self.assertIn("(plate_readout)", assistant_tools.INSTRUCTIONS)
        local = (Path(__file__).resolve().parent.parent / "mcp" / "biomanager_mcp.py").read_text(encoding="utf-8")
        self.assertIn('@server.tool(name="plate_readout"', local)


if __name__ == "__main__":
    unittest.main()


class DeletingTheExperiment(AppTestCase):
    def test_the_colony_s_delete_takes_its_readings_too(self):
        # Readings (a plate's among them) aren't the experiment's ORM
        # children; left behind, PostgreSQL refuses the delete.
        r = self.m.post("/colony/experiments/create", data={"name": uniq("Plate exp ")})
        exp = int(location(r).rsplit("/", 1)[1])
        from tests.base import execute
        execute("insert into experiment_readings (experiment_id_fk, subject, readout_key, read_on, value, note, "
                "recorded_by, created_at) values (?, 'mouse:1', 'custom', ?, 0.5, '', ?, ?)",
                exp, TODAY, self.member, TODAY)
        self.m.post(f"/colony/experiments/{exp}/delete")
        self.assertEqual(count("experiments", "id=?", exp), 0)
        self.assertEqual(count("experiment_readings", "experiment_id_fk=?", exp), 0)
